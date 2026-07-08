# -*- coding: utf-8 -*-
"""Excel writer / merger for the BOQ export tools.

Uses the vendored openpyxl (lib/vendor) so no Microsoft Excel is required.

Merge contract
--------------
* Rows are matched on the ``Key`` column (Revit UniqueId, or
  ``<typeUID>::L<n>`` for layer rows).
* Plugin-owned columns (those the extractor produces) are refreshed on every
  export; any *extra* columns a user added (QS rates, prices, notes) are kept
  untouched.
* The ``Status`` column is recomputed each run:
  New / Changed / Unchanged / Removed. Removed rows are kept (so pricing is not
  lost) and flagged rather than deleted.
"""

import os
import sys

# make the vendored openpyxl importable (.../lib/bkbim/boq/xlio.py -> .../lib/vendor)
_VENDOR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "vendor")
if _VENDOR not in sys.path:
    sys.path.insert(0, _VENDOR)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from bkbim.boq.model import KEY, STATUS

# ---- styling --------------------------------------------------------------
FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3B4D")      # dark slate
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {
    "New": PatternFill("solid", fgColor="C6EFCE"),        # green
    "Changed": PatternFill("solid", fgColor="FFEB9C"),    # amber
    "Removed": PatternFill("solid", fgColor="FFC7CE"),    # red
    "Unchanged": None,
}

NUMERIC_HINTS = ("(mm)", "(m2)", "(m3)", "(m)", "Count")


def _num_format(header):
    if "(mm)" in header:
        return "#,##0"
    if "(kg)" in header:
        return "#,##0.0"
    if "(m2)" in header or "(m3)" in header:
        return "#,##0.00"
    if "(m)" in header:
        return "#,##0.000"
    if header == "Count":
        return "#,##0"
    return None


def _safe_sheet_name(name):
    for ch in "[]:*?/\\":
        name = name.replace(ch, "-")
    return name[:31]


def _norm(v):
    """Normalise a value for change comparison."""
    if v is None:
        return ""
    if isinstance(v, float):
        return round(v, 4)
    if isinstance(v, int):
        return float(v) if v != int(v) else v
    s = str(v).strip()
    return s


def _eq(a, b):
    na, nb = _norm(a), _norm(b)
    try:
        return abs(float(na) - float(nb)) < 1e-6
    except (TypeError, ValueError):
        return str(na) == str(nb)


def _read_sheet(ws):
    """Return (header_list, {key: row_dict}) from an existing worksheet."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {}
    headers = [str(h) if h is not None else "" for h in rows[0]]
    by_key = {}
    try:
        kidx = headers.index(KEY)
    except ValueError:
        return headers, {}
    for r in rows[1:]:
        if kidx >= len(r) or r[kidx] is None:
            continue
        d = {}
        for i, h in enumerate(headers):
            d[h] = r[i] if i < len(r) else None
        by_key[str(r[kidx])] = d
    return headers, by_key


def _merge_rows(plugin_cols, new_rows, existing_headers, existing_by_key):
    """Return (final_columns, final_rows) after merging new_rows into existing."""
    extra_cols = [h for h in existing_headers
                  if h not in plugin_cols and h not in (KEY, STATUS)]
    final_cols = list(plugin_cols) + extra_cols
    data_cols = [c for c in plugin_cols if c not in (KEY, STATUS)]

    final_rows = []
    seen = set()
    for row in new_rows:
        key = str(row.get(KEY))
        seen.add(key)
        old = existing_by_key.get(key)
        merged = {}
        if old is None:
            for c in final_cols:
                merged[c] = row.get(c, "")
            merged[STATUS] = "New"
        else:
            changed = any(not _eq(old.get(c), row.get(c)) for c in data_cols)
            for c in final_cols:
                if c in plugin_cols:
                    merged[c] = row.get(c, "")
                else:
                    merged[c] = old.get(c, "")   # preserve QS column
            merged[STATUS] = "Changed" if changed else "Unchanged"
        final_rows.append(merged)

    # keep rows that vanished from the model
    for key, old in existing_by_key.items():
        if key in seen:
            continue
        merged = {c: old.get(c, "") for c in final_cols}
        merged[KEY] = key
        merged[STATUS] = "Removed"
        final_rows.append(merged)

    return final_cols, final_rows


def _fresh_rows(plugin_cols, new_rows):
    final_rows = []
    for row in new_rows:
        merged = {c: row.get(c, "") for c in plugin_cols}
        merged[STATUS] = "New"
        final_rows.append(merged)
    return list(plugin_cols), final_rows


def _write_ws(ws, columns, rows):
    # header
    for ci, header in enumerate(columns, start=1):
        c = ws.cell(row=1, column=ci, value=header)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    # data
    status_idx = columns.index(STATUS) + 1 if STATUS in columns else None
    widths = [len(str(h)) for h in columns]
    for ri, row in enumerate(rows, start=2):
        for ci, header in enumerate(columns, start=1):
            val = row.get(header, "")
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = DATA_FONT
            cell.border = BORDER
            fmt = _num_format(header)
            if fmt and isinstance(val, (int, float)):
                cell.number_format = fmt
            L = len(str(val)) if val is not None else 0
            if L > widths[ci - 1]:
                widths[ci - 1] = L
        if status_idx:
            st = row.get(STATUS, "")
            fill = STATUS_FILL.get(st)
            if fill:
                ws.cell(row=ri, column=status_idx).fill = fill
    # widths / freeze / filter
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = min(max(w + 2, 8), 55)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(columns)), len(rows) + 1)


def write_extracts(path, extracts, merge_mode):
    """Write/merge a list of Extract objects to ``path``.

    merge_mode: "merge" to update an existing workbook (preserving extra
    columns), or "new" to start a clean workbook.
    """
    do_merge = merge_mode == "merge" and os.path.exists(path)
    wb = openpyxl.load_workbook(path) if do_merge else openpyxl.Workbook()
    if not do_merge:
        # drop the default empty sheet
        for s in list(wb.sheetnames):
            del wb[s]

    summary = []
    for ex in extracts:
        for sheet in ex.sheets:
            sheet_name = _safe_sheet_name("{} - {}".format(ex.label, sheet.name))
            if do_merge and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                headers, by_key = _read_sheet(ws)
                cols, rows = _merge_rows(sheet.columns, sheet.rows, headers, by_key)
                wb.remove(ws)
            else:
                cols, rows = _fresh_rows(sheet.columns, sheet.rows)
            ws = wb.create_sheet(title=sheet_name)
            _write_ws(ws, cols, rows)
            counts = {}
            for r in rows:
                counts[r.get(STATUS, "")] = counts.get(r.get(STATUS, ""), 0) + 1
            summary.append((sheet_name, len(rows), counts))

    folder = os.path.dirname(path)
    if folder and not os.path.exists(folder):
        os.makedirs(folder)
    wb.save(path)
    return summary


def split_workbook(path, out_folder):
    """Explode every sheet of a workbook into its own single-sheet .xlsx."""
    wb = openpyxl.load_workbook(path)
    written = []
    base = os.path.splitext(os.path.basename(path))[0]
    if not os.path.exists(out_folder):
        os.makedirs(out_folder)
    for name in wb.sheetnames:
        src = wb[name]
        nwb = openpyxl.Workbook()
        del nwb[nwb.sheetnames[0]]
        nws = nwb.create_sheet(title=_safe_sheet_name(name))
        for row in src.iter_rows():
            for cell in row:
                nc = nws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    nc.font = cell.font.copy()
                    nc.fill = cell.fill.copy()
                    nc.border = cell.border.copy()
                    nc.alignment = cell.alignment.copy()
                    nc.number_format = cell.number_format
        for col, dim in src.column_dimensions.items():
            nws.column_dimensions[col].width = dim.width
        nws.freeze_panes = "A2"
        safe = name.replace(" ", "_").replace("/", "-")
        out = os.path.join(out_folder, "{}__{}.xlsx".format(base, safe))
        nwb.save(out)
        written.append(out)
    return written
